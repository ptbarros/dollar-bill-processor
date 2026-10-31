#!/usr/bin/env python3
"""
Apply curated pattern aliases from the review spreadsheet into the .lua AKA headers.

Reads the alias worksheet (the one produced for the AKA feature: one row per core
pattern, with a "Filename" column, an "Internal Name" column, and a
"Proposed Aliases" column), and writes each row's aliases into that pattern's Lua
header as an ``AKA:`` line. These are the SHIPPED aliases that travel with the
pattern (and get vendored to dd-web) -- distinct from the per-user alias overrides
stored in settings.

This is a maintenance helper, NOT wired into the app.

    # see what would change, without touching anything (default):
    python tools/apply_pattern_aliases.py pattern_aliases.xlsx

    # actually write the AKA headers:
    python tools/apply_pattern_aliases.py pattern_aliases.xlsx --write

Idempotent: an existing ``AKA:`` line is replaced in place; otherwise a new one is
inserted after Tier (or Examples/Description) in the header block. A row with an
empty "Proposed Aliases" cell is left alone (its existing header, if any, is kept)
unless --prune is given, which removes the AKA line for empty rows.

By default the four names still under naming review with Ed are SKIPPED
(adding an alias is harmless, but this keeps the change off those rows until the
display names are settled). Pass --include-review to apply them too.
"""

import argparse
import re
import sys
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parent.parent
CORE_DIR = REPO_ROOT / "patterns" / "core"

# Names whose DISPLAY name is still under review with Ed (see project memory).
# Aliases are safe to add, but we hold off by default to avoid churn on them.
UNDER_REVIEW = {
    "TRUE_BINARY_SKIP_NOTE", "TRUE_DOUBLE_QUAD_BINARY",
    "RANDOM_DOUBLE_QUAD_BINARY", "TETRADIC",
}

# Column header -> normalized key. We match leniently so a lightly-renamed column
# (e.g. "Proposed Aliases (edit me)") still resolves.
def _find_columns(header_row):
    cols = {}
    for idx, cell in enumerate(header_row):
        text = (str(cell.value) if cell.value is not None else "").strip().lower()
        if not text:
            continue
        if text.startswith("filename"):
            cols["filename"] = idx
        elif text.startswith("internal"):
            cols["name"] = idx
        elif "alias" in text and "existing" not in text:
            cols.setdefault("aliases", idx)  # first alias-ish col that isn't "existing"
    return cols


def _parse_aliases(cell_value):
    if cell_value is None:
        return []
    return [a.strip() for a in str(cell_value).split(",") if a.strip()]


def _set_aka_in_header(text, aliases, prune=False):
    """Return (new_text, action) with the AKA line set/replaced/removed inside the
    first --[[ ... --]] header block. action is 'replaced', 'inserted', 'removed',
    or 'none'. Raises ValueError if there is no block header."""
    m = re.search(r"(--\[\[)(.*?)(--\]\])", text, re.DOTALL)
    if not m:
        raise ValueError("no --[[ ... --]] header block")
    head_open, body, head_close = m.group(1), m.group(2), m.group(3)
    lines = body.split("\n")

    aka_line = "AKA: " + ", ".join(aliases) if aliases else ""
    existing_idx = next(
        (i for i, ln in enumerate(lines) if re.match(r"\s*AKA\s*:", ln, re.I)), None)

    if existing_idx is not None:
        if aliases:
            if lines[existing_idx].strip() == aka_line:
                return text, "none"
            lines[existing_idx] = aka_line
            action = "replaced"
        elif prune:
            lines.pop(existing_idx)
            action = "removed"
        else:
            return text, "none"
    else:
        if not aliases:
            return text, "none"
        # Insert after Tier, else Examples, else Description, else before close.
        anchor = None
        for key in ("tier", "examples", "description"):
            anchor = next(
                (i for i, ln in enumerate(lines) if re.match(rf"\s*{key}\s*:", ln, re.I)),
                None)
            if anchor is not None:
                break
        insert_at = (anchor + 1) if anchor is not None else len(lines) - 1
        lines.insert(insert_at, aka_line)
        action = "inserted"

    new_body = "\n".join(lines)
    new_text = text[:m.start()] + head_open + new_body + head_close + text[m.end():]
    return new_text, action


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("workbook", help="the curated alias .xlsx")
    ap.add_argument("--write", action="store_true",
                    help="apply the changes (default: dry run, show what would change)")
    ap.add_argument("--prune", action="store_true",
                    help="also remove the AKA line for rows whose alias cell is empty")
    ap.add_argument("--include-review", action="store_true",
                    help="also apply the under-review names (skipped by default)")
    ap.add_argument("--sheet", default=None, help="worksheet name (default: first)")
    args = ap.parse_args()

    try:
        import openpyxl
    except ImportError:
        sys.exit("openpyxl is required (use the project venv: venv/bin/python ...).")

    wb = openpyxl.load_workbook(args.workbook, read_only=True, data_only=True)
    ws = wb[args.sheet] if args.sheet else wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows())
    if not rows:
        sys.exit("empty worksheet")

    cols = _find_columns(rows[0])
    for need in ("filename", "aliases"):
        if need not in cols:
            sys.exit(f"could not find a '{need}' column in the header row; "
                     f"found columns: {[c.value for c in rows[0]]}")
    name_col = cols.get("name")

    changed = skipped_review = unchanged = blank = errors = 0
    actions = []
    for r in rows[1:]:
        fname = r[cols["filename"]].value
        if not fname:
            continue
        fname = str(fname).strip()
        name = str(r[name_col].value).strip() if name_col is not None and r[name_col].value else fname
        aliases = _parse_aliases(r[cols["aliases"]].value)

        if name in UNDER_REVIEW and not args.include_review:
            skipped_review += 1
            continue
        if not aliases and not args.prune:
            blank += 1
            continue

        path = CORE_DIR / fname
        if not path.exists():
            actions.append(("ERROR", fname, f"file not found: {path}"))
            errors += 1
            continue
        text = path.read_text(encoding="utf-8")
        try:
            new_text, action = _set_aka_in_header(text, aliases, prune=args.prune)
        except ValueError as e:
            actions.append(("ERROR", fname, str(e)))
            errors += 1
            continue

        if action == "none":
            unchanged += 1
            continue
        actions.append((action, fname, ", ".join(aliases) if aliases else "(removed)"))
        changed += 1
        if args.write:
            path.write_text(new_text, encoding="utf-8")

    for action, fname, detail in actions:
        mark = {"inserted": "+", "replaced": "~", "removed": "-", "ERROR": "!"}.get(action, "?")
        print(f"  {mark} {fname:34} {detail}")

    mode = "WROTE" if args.write else "DRY RUN (no files changed) --"
    print(f"\n{mode}  changed: {changed}  unchanged: {unchanged}  "
          f"blank-skipped: {blank}  review-skipped: {skipped_review}  errors: {errors}")
    if not args.write and changed:
        print("Re-run with --write to apply.")
    if errors:
        sys.exit(1)


if __name__ == "__main__":
    main()
